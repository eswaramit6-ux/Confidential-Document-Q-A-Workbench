"""
Text extraction for PDF and DOCX files. Runs entirely locally using
pypdf and python-docx — no external document-processing API is used.
"""
import hashlib
from dataclasses import dataclass, field
from typing import List


class UnsupportedFileTypeError(ValueError):
    pass


class CorruptedDocumentError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


@dataclass
class PageText:
    page_number: int  # 1-indexed; None-like sentinel of 0 used when not applicable
    text: str


@dataclass
class ExtractedDocument:
    filename: str
    sha256: str
    pages: List[PageText] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.pages)

    @property
    def has_page_numbers(self) -> bool:
        return any(p.page_number > 0 for p in self.pages)


def compute_sha256(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def extract_pdf(file_bytes: bytes, filename: str) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(file_bytes))
        if len(reader.pages) == 0:
            raise EmptyDocumentError(f"'{filename}' has no pages.")
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            pages.append(PageText(page_number=i, text=text))
    except EmptyDocumentError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise CorruptedDocumentError(f"Could not read PDF '{filename}': {exc}") from exc

    if not any(p.text.strip() for p in pages):
        raise EmptyDocumentError(
            f"'{filename}' appears to contain no extractable text (it may be a scanned image PDF)."
        )
    return ExtractedDocument(filename=filename, sha256=compute_sha256(file_bytes), pages=pages)


def extract_docx(file_bytes: bytes, filename: str) -> ExtractedDocument:
    try:
        import docx
        import io
        document = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        text = "\n".join(paragraphs)
    except Exception as exc:  # noqa: BLE001
        raise CorruptedDocumentError(f"Could not read DOCX '{filename}': {exc}") from exc

    if not text.strip():
        raise EmptyDocumentError(f"'{filename}' appears to contain no extractable text.")

    # DOCX has no reliable page concept; store as a single logical "page" (page_number=0 => N/A)
    return ExtractedDocument(
        filename=filename,
        sha256=compute_sha256(file_bytes),
        pages=[PageText(page_number=0, text=text)],
    )


def extract_document(file_bytes: bytes, filename: str) -> ExtractedDocument:
    lower = filename.lower()
    if not file_bytes:
        raise EmptyDocumentError(f"'{filename}' is an empty file.")
    if lower.endswith(".pdf"):
        return extract_pdf(file_bytes, filename)
    if lower.endswith(".docx"):
        return extract_docx(file_bytes, filename)
    raise UnsupportedFileTypeError(f"Unsupported file type for '{filename}'. Only PDF and DOCX are supported.")
