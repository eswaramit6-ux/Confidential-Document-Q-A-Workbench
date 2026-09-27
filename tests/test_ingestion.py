import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

from ingestion.chunker import chunk_document
from ingestion.document_loader import (
    EmptyDocumentError,
    UnsupportedFileTypeError,
    compute_sha256,
    extract_document,
)


def test_docx_extraction_and_sha256():
    import docx
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("This is a confidential test policy document.")
    document.add_paragraph("It contains an emergency shutdown procedure.")
    document.save(buf)
    file_bytes = buf.getvalue()

    extracted = extract_document(file_bytes, "policy.docx")
    assert "emergency shutdown procedure" in extracted.full_text
    assert extracted.sha256 == compute_sha256(file_bytes)
    assert not extracted.has_page_numbers


def test_unsupported_file_type():
    with pytest.raises(UnsupportedFileTypeError):
        extract_document(b"hello", "notes.txt")


def test_empty_docx_raises():
    import docx
    buf = io.BytesIO()
    document = docx.Document()
    document.save(buf)
    with pytest.raises(EmptyDocumentError):
        extract_document(buf.getvalue(), "empty.docx")


def test_chunking_produces_overlapping_chunks():
    import docx
    buf = io.BytesIO()
    document = docx.Document()
    long_text = " ".join([f"Sentence number {i} about safety procedures." for i in range(200)])
    document.add_paragraph(long_text)
    document.save(buf)
    extracted = extract_document(buf.getvalue(), "long.docx")

    chunks = chunk_document(extracted, chunk_size=300, overlap=50)
    assert len(chunks) > 1
    assert all(c.filename == "long.docx" for c in chunks)
    assert all(c.sha256 == extracted.sha256 for c in chunks)


def test_duplicate_sha256_detection_logic():
    content = b"identical file bytes for duplicate test"
    sha_a = compute_sha256(content)
    sha_b = compute_sha256(content)
    assert sha_a == sha_b
